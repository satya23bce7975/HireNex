"""
HireNex - Hybrid Retrieval / RAG Engine

Purpose
-------
This module is responsible ONLY for retrieval, candidate filtering,
semantic search, grounded Q&A, and recruiter-result Excel generation.

Input sources
-------------
1. output/nlp/resumes.json
   Traditional NLP structured data.
2. output/llm/resumes.json
   LLM structured data.
3. md_data/*.md
   Canonical Markdown resume documents for semantic retrieval.
4. data/resumes/*.pdf
   Original PDFs used as source evidence / verification.

Architecture
------------
                         Recruiter Query
                               |
                               v
                       Query Understanding
                               |
               +---------------+---------------+
               |                               |
               v                               v
        Structured filters               Semantic query
               |                               |
               v                               v
       NLP JSON + LLM JSON              ChromaDB
               |                       MD + PDF chunks
               +---------------+---------------+
                               |
                               v
                     Candidate ranking/results
                               |
                 +-------------+-------------+
                 |                           |
                 v                           v
             Answer                    Excel export

Important:
- Exact filters (experience, skill, CGPA, degree, college, etc.)
  are applied deterministically in Python.
- Semantic questions use vector retrieval.
- Mixed questions can use BOTH paths.
- The original PDF is retained as evidence/source verification.
"""

from __future__ import annotations

import json
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from groq import Groq
import chromadb
from sentence_transformers import SentenceTransformer

# PDF support is used only for source/evidence extraction.
try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

# Excel generation uses openpyxl so the project does not depend on artifact_tool.
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:
    Workbook = None
    Font = None
    PatternFill = None
    Alignment = None
    Table = None
    TableStyleInfo = None


# =============================================================================
# CONFIGURATION
# =============================================================================

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent

NLP_JSON_PATH = BASE_DIR / os.getenv(
    "NLP_JSON_PATH",
    "output/nlp/resumes.json",
)

LLM_JSON_PATH = BASE_DIR / os.getenv(
    "LLM_JSON_PATH",
    "output/llm/resumes.json",
)

MD_FOLDER = BASE_DIR / os.getenv(
    "MD_FOLDER",
    "md_data",
)

PDF_FOLDER = BASE_DIR / os.getenv(
    "PDF_FOLDER",
    "data/resumes",
)

RAG_FOLDER = BASE_DIR / os.getenv(
    "RAG_FOLDER",
    "rag",
)

CHROMA_FOLDER = RAG_FOLDER / "chroma_db"

RAG_OUTPUT_FOLDER = BASE_DIR / os.getenv(
    "RAG_OUTPUT_FOLDER",
    "ragoutput",
)

INDEX_META_PATH = RAG_FOLDER / "index_meta.json"

RAG_FOLDER.mkdir(parents=True, exist_ok=True)
RAG_OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)


GROQ_API_KEY = os.getenv("GROQ_API_KEY", "").strip()

if not GROQ_API_KEY:
    raise ValueError(
        "GROQ_API_KEY is missing. Add it to your .env file."
    )

client = Groq(api_key=GROQ_API_KEY)

QUERY_MODEL = os.getenv(
    "GROQ_QUERY_MODEL",
    "openai/gpt-oss-120b",
)

ANSWER_MODEL = os.getenv(
    "GROQ_ANSWER_MODEL",
    "openai/gpt-oss-120b",
)

EMBEDDING_MODEL = os.getenv(
    "EMBEDDING_MODEL",
    "sentence-transformers/all-MiniLM-L6-v2",
)

# Keep chunks reasonably small for retrieval quality.
MD_CHUNK_LENGTH = 1200
PDF_CHUNK_LENGTH = 1200
CHUNK_OVERLAP = 150

DEFAULT_SEMANTIC_RESULTS = 20
DEFAULT_CONTEXT_RESULTS = 8


# =============================================================================
# CHROMA + EMBEDDINGS
# =============================================================================

embedding_model = SentenceTransformer(EMBEDDING_MODEL)

chroma_client = chromadb.PersistentClient(
    path=str(CHROMA_FOLDER)
)

collection = chroma_client.get_or_create_collection(
    name="hinenex_resume_knowledge",
    metadata={"hnsw:space": "cosine"},
)


# =============================================================================
# GENERIC HELPERS
# =============================================================================

def clean_text(value: Any) -> str:
    if value is None:
        return ""

    text = str(value)
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)

    return text.strip()


def normalize_key(value: Any) -> str:
    return clean_text(value).lower()


def dedupe_list(values: list[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()

    for value in values:
        value = clean_text(value)

        if not value:
            continue

        key = value.lower()

        if key not in seen:
            seen.add(key)
            result.append(value)

    return result


def unique_json_objects(
    objects: list[dict[str, Any]],
    keys: tuple[str, ...],
) -> list[dict[str, Any]]:

    result: list[dict[str, Any]] = []
    seen: set[str] = set()

    for obj in objects:

        signature = "|".join(
            normalize_key(obj.get(key, ""))
            for key in keys
        )

        if not signature.strip("|"):
            continue

        if signature not in seen:

            seen.add(signature)
            result.append(obj)

    return result


def ensure_list(value: Any) -> list[Any]:

    if value is None:
        return []

    if isinstance(value, list):
        return value

    return [value]


# =============================================================================
# JSON LOADING
# =============================================================================

def load_json_file(path: Path) -> list[dict[str, Any]]:

    if not path.exists():
        return []

    with path.open(
        "r",
        encoding="utf-8",
    ) as file:

        data = json.load(file)

    if isinstance(data, dict):

        if isinstance(data.get("resumes"), list):
            data = data["resumes"]

        elif isinstance(data.get("candidates"), list):
            data = data["candidates"]

        else:
            data = [data]

    if not isinstance(data, list):
        raise ValueError(
            f"{path} must contain a JSON list or a supported wrapper."
        )

    return [
        item
        for item in data
        if isinstance(item, dict)
    ]


def load_source_jsons() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:

    nlp_records = load_json_file(
        NLP_JSON_PATH
    )

    llm_records = load_json_file(
        LLM_JSON_PATH
    )

    return nlp_records, llm_records


# =============================================================================
# CANDIDATE NORMALIZATION
# =============================================================================

def get_application_no(record: dict[str, Any]) -> str:

    value = clean_text(
        record.get("application_no", "")
    )

    if value:
        return value.upper()

    metadata = record.get(
        "metadata",
        {}
    )

    if isinstance(metadata, dict):

        value = clean_text(
            metadata.get("application_no", "")
        )

        if value:
            return value.upper()

        for field in (
            "source_markdown",
            "file_name",
            "source_file",
        ):

            source = clean_text(
                metadata.get(field, "")
            )

            if source:

                stem = Path(
                    source
                ).stem.upper()

                if stem.startswith("APP"):
                    return stem

    return ""


def candidate_block(
    record: dict[str, Any],
) -> dict[str, Any]:

    block = record.get(
        "candidate",
        {}
    )

    if not isinstance(block, dict):
        block = {}

    return block


def contact_block(
    record: dict[str, Any],
) -> dict[str, Any]:

    candidate = candidate_block(record)

    contact = candidate.get(
        "contact",
        {}
    )

    if not isinstance(contact, dict):
        contact = {}

    return contact


def merge_scalar(
    preferred: Any,
    fallback: Any,
) -> str:

    preferred_text = clean_text(
        preferred
    )

    if preferred_text:
        return preferred_text

    return clean_text(
        fallback
    )


def normalize_experience_item(
    item: dict[str, Any],
) -> dict[str, Any]:

    # NLP schema
    role = clean_text(
        item.get(
            "role",
            item.get("title", "")
        )
    )

    company = clean_text(
        item.get(
            "company",
            item.get("Company", "")
        )
    )

    description = clean_text(
        item.get(
            "description",
            ""
        )
    )

    duration = clean_text(
        item.get(
            "duration",
            ""
        )
    )

    responsibilities = item.get(
        "responsibilities",
        [],
    )

    if responsibilities:

        responsibilities_text = " ".join(
            clean_text(x)
            for x in ensure_list(responsibilities)
        )

        if responsibilities_text:
            description = merge_scalar(
                description,
                responsibilities_text,
            )

    start_year = clean_text(
        item.get(
            "start_year",
            ""
        )
    )

    end_year = clean_text(
        item.get(
            "end_year",
            ""
        )
    )

    return {
        "company": company,
        "role": role,
        "description": description,
        "duration": duration,
        "start_year": start_year,
        "end_year": end_year,
    }


def merge_candidates(
    nlp_records: list[dict[str, Any]],
    llm_records: list[dict[str, Any]],
) -> list[dict[str, Any]]:

    by_app: dict[str, dict[str, Any]] = {}

    # Put NLP first so every candidate is represented even when
    # the LLM parser missed a record.
    for source_name, records in (
        ("nlp", nlp_records),
        ("llm", llm_records),
    ):

        for record in records:

            app_no = get_application_no(
                record
            )

            if not app_no:
                continue

            if app_no not in by_app:

                by_app[app_no] = {
                    "_nlp": {},
                    "_llm": {},
                }

            by_app[app_no][
                f"_{source_name}"
            ] = record

    merged: list[dict[str, Any]] = []

    for app_no in sorted(by_app):

        nlp = by_app[app_no].get(
            "_nlp",
            {}
        )

        llm = by_app[app_no].get(
            "_llm",
            {}
        )

        nlp_candidate = candidate_block(nlp)
        llm_candidate = candidate_block(llm)

        nlp_contact = contact_block(nlp)
        llm_contact = contact_block(llm)

        # LLM is preferred where it has a value;
        # NLP becomes the deterministic fallback.
        merged_candidate = {

            "application_no": app_no,

            "name": merge_scalar(
                llm_candidate.get("name"),
                nlp_candidate.get("name"),
            ),

            "tag": merge_scalar(
                llm_candidate.get("tag"),
                nlp_candidate.get("tag"),
            ),

            "email": merge_scalar(
                llm_contact.get("email"),
                nlp_contact.get("email"),
            ),

            "phone": merge_scalar(
                llm_contact.get("phone"),
                nlp_contact.get("phone"),
            ),

            "linkedin": merge_scalar(
                llm_contact.get("linkedin"),
                nlp_contact.get("linkedin"),
            ),

            "github": merge_scalar(
                llm_contact.get("github"),
                nlp_contact.get("github"),
            ),

            "portfolio": merge_scalar(
                llm_contact.get("portfolio"),
                nlp_contact.get("portfolio"),
            ),

            "personal_summary": merge_scalar(
                llm_candidate.get("personal_summary"),
                nlp_candidate.get("personal_summary"),
            ),

            "skills": dedupe_list(
                [
                    *ensure_list(
                        nlp.get("skills", [])
                    ),
                    *ensure_list(
                        llm.get("skills", [])
                    ),
                ]
            ),

            "education": unique_json_objects(
                [
                    *[
                        x for x in ensure_list(
                            nlp.get("education", [])
                        )
                        if isinstance(x, dict)
                    ],
                    *[
                        x for x in ensure_list(
                            llm.get("education", [])
                        )
                        if isinstance(x, dict)
                    ],
                ],
                ("degree", "field", "institute"),
            ),

            "projects": unique_json_objects(
                [
                    *[
                        x for x in ensure_list(
                            nlp.get("projects", [])
                        )
                        if isinstance(x, dict)
                    ],
                    *[
                        x for x in ensure_list(
                            llm.get("projects", [])
                        )
                        if isinstance(x, dict)
                    ],
                ],
                ("name",),
            ),

            "experience": merge_experience_items(
                [
                    *[
                        normalize_experience_item(x)
                        for x in ensure_list(
                            nlp.get("experience", [])
                        )
                        if isinstance(x, dict)
                    ],
                    *[
                        normalize_experience_item(x)
                        for x in ensure_list(
                            llm.get("experience", [])
                        )
                        if isinstance(x, dict)
                    ],
                ]
            ),

            "certifications": unique_json_objects(
                [
                    *[
                        x for x in ensure_list(
                            nlp.get("certifications", [])
                        )
                        if isinstance(x, dict)
                    ],
                    *[
                        x for x in ensure_list(
                            llm.get("certifications", [])
                        )
                        if isinstance(x, dict)
                    ],
                ],
                ("name", "issuer", "year"),
            ),

            "achievements": dedupe_list(
                [
                    *ensure_list(
                        nlp.get("achievements", [])
                    ),
                    *ensure_list(
                        llm.get("achievements", [])
                    ),
                ]
            ),
        }

        merged_candidate[
            "total_experience_years"
        ] = calculate_total_experience(
            merged_candidate["experience"]
        )

        merged_candidate[
            "max_cgpa"
        ] = extract_best_cgpa(
            merged_candidate["education"]
        )

        merged_candidate[
            "max_percentage"
        ] = extract_best_percentage(
            merged_candidate["education"]
        )

        merged_candidate[
            "source_available"
        ] = {
            "nlp": bool(nlp),
            "llm": bool(llm),
            "markdown": (
                find_markdown(app_no) is not None
            ),
            "pdf": (
                find_pdf(app_no) is not None
            ),
        }

        merged.append(
            merged_candidate
        )

    return merged


def merge_experience_items(
    items: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Merge NLP/LLM entries representing the same job."""
    grouped: dict[str, dict[str, Any]] = {}

    for item in items:
        company = clean_text(item.get("company", ""))
        role = clean_text(item.get("role", ""))
        key = f"{normalize_key(company)}|{normalize_key(role)}"

        if not key.strip("|"):
            # Preserve unusual entries instead of silently losing them.
            key = f"__unique__{len(grouped)}"

        current = grouped.get(key)

        if current is None:
            grouped[key] = {
                "company": company,
                "role": role,
                "description": clean_text(item.get("description", "")),
                "duration": clean_text(item.get("duration", "")),
                "start_year": clean_text(item.get("start_year", "")),
                "end_year": clean_text(item.get("end_year", "")),
            }
            continue

        current["company"] = merge_scalar(
            current.get("company"), company
        )
        current["role"] = merge_scalar(
            current.get("role"), role
        )
        current["description"] = merge_scalar(
            current.get("description"),
            clean_text(item.get("description", "")),
        )
        current["duration"] = merge_scalar(
            current.get("duration"),
            clean_text(item.get("duration", "")),
        )
        current["start_year"] = merge_scalar(
            current.get("start_year"),
            clean_text(item.get("start_year", "")),
        )
        current["end_year"] = merge_scalar(
            current.get("end_year"),
            clean_text(item.get("end_year", "")),
        )

    return list(grouped.values())


# =============================================================================
# DATE / EXPERIENCE CALCULATION
# =============================================================================

MONTHS = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}


def parse_month_year(
    text: str,
) -> tuple[int, int] | None:

    match = re.search(
        r"\b("
        r"jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|"
        r"jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|"
        r"oct(?:ober)?|nov(?:ember)?|dec(?:ember)?"
        r")\.?\s+"
        r"((?:19|20)\d{2})\b",
        text,
        re.I,
    )

    if not match:
        return None

    month_text = match.group(1).lower()
    year = int(match.group(2))

    if month_text.startswith("sep"):
        month = 9
    else:
        month = MONTHS.get(
            month_text[:3],
            1,
        )

    return year, month


def parse_year(
    text: str,
) -> int | None:

    match = re.search(
        r"\b((?:19|20)\d{2})\b",
        text,
    )

    if not match:
        return None

    return int(
        match.group(1)
    )


def duration_to_months(
    duration: str,
) -> int:

    text = clean_text(
        duration
    ).lower()

    if not text:
        return 0

    # Explicit "2 years 3 months"
    year_match = re.search(
        r"(\d+(?:\.\d+)?)\s+years?",
        text,
        re.I,
    )

    month_match = re.search(
        r"(\d+)\s+months?",
        text,
        re.I,
    )

    if year_match:

        months = round(
            float(year_match.group(1))
            * 12
        )

        if month_match:
            months += int(
                month_match.group(1)
            )

        return months

    # Month/year range, e.g. Jan 2022 - Dec 2024
    dates = re.findall(
        r"\b(?:"
        r"jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|"
        r"jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|"
        r"oct(?:ober)?|nov(?:ember)?|dec(?:ember)?"
        r")\.?\s+(?:19|20)\d{2}\b",
        text,
        re.I,
    )

    if len(dates) >= 2:

        start = parse_month_year(
            dates[0]
        )

        end = parse_month_year(
            dates[1]
        )

        if start and end:

            months = (
                (end[0] - start[0]) * 12
                + (end[1] - start[1])
            )

            return max(
                months,
                0,
            )

    # Year-only range, e.g. 2022 - 2024
    year_range = re.search(
        r"\b((?:19|20)\d{2})\s*"
        r"(?:-|–|—|to)\s*"
        r"((?:19|20)\d{2}|present|current)\b",
        text,
        re.I,
    )

    if year_range:

        start_year = int(
            year_range.group(1)
        )

        end_value = year_range.group(2)

        if end_value.lower() in {
            "present",
            "current",
        }:
            end_year = datetime.now().year

        else:
            end_year = int(
                end_value
            )

        return max(
            (end_year - start_year) * 12,
            0,
        )

    return 0


def experience_item_to_months(
    experience: dict[str, Any],
) -> int:

    duration = clean_text(
        experience.get("duration", "")
    )

    months = duration_to_months(
        duration
    )

    if months:
        return months

    start = parse_year(
        clean_text(
            experience.get("start_year", "")
        )
    )

    end_text = clean_text(
        experience.get("end_year", "")
    )

    if start is None:
        return 0

    if end_text.lower() in {
        "present",
        "current",
        "now",
    }:

        end = datetime.now().year

    else:

        end = parse_year(
            end_text
        )

    if end is None:
        return 0

    return max(
        (end - start) * 12,
        0,
    )


def experience_item_interval(
    experience: dict[str, Any],
) -> tuple[int, int] | None:
    """Return an approximate inclusive month interval for one job."""
    duration = clean_text(experience.get("duration", ""))

    # Month-year interval from the LLM duration field.
    month_matches = re.findall(
        r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|"
        r"jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|"
        r"oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)\.?\s+(?:19|20)\d{2}",
        duration,
        re.I,
    )

    if len(month_matches) >= 2:
        start = parse_month_year(month_matches[0])
        end = parse_month_year(month_matches[1])
        if start and end:
            return (
                start[0] * 12 + start[1],
                end[0] * 12 + end[1],
            )

    # Year-only interval from either parser.
    start_year = parse_year(
        clean_text(experience.get("start_year", ""))
    )
    end_text = clean_text(
        experience.get("end_year", "")
    )

    if start_year is None:
        return None

    if end_text.lower() in {"present", "current", "now"}:
        end_year = datetime.now().year
    else:
        end_year = parse_year(end_text)

    if end_year is None:
        return None

    return (
        start_year * 12 + 1,
        end_year * 12 + 12,
    )


def calculate_total_experience(
    experiences: list[dict[str, Any]],
) -> float:
    """Calculate approximate total experience without double-counting overlap."""
    intervals: list[tuple[int, int]] = []

    for experience in experiences:
        interval = experience_item_interval(experience)
        if interval:
            intervals.append(interval)

    if intervals:
        intervals.sort()
        merged: list[list[int]] = []

        for start, end in intervals:
            if not merged or start > merged[-1][1] + 1:
                merged.append([start, end])
            else:
                merged[-1][1] = max(merged[-1][1], end)

        total_months = sum(
            end - start + 1
            for start, end in merged
        )

        return round(total_months / 12, 1)

    # If dates are not parseable, fall back to explicit durations.
    total_months = sum(
        duration_to_months(
            clean_text(experience.get("duration", ""))
        )
        for experience in experiences
    )

    return round(total_months / 12, 1)


def extract_best_cgpa(
    education: list[dict[str, Any]],
) -> float | None:

    values: list[float] = []

    for item in education:

        score = item.get(
            "score",
            {}
        )

        if not isinstance(score, dict):
            continue

        score_type = normalize_key(
            score.get("type", "")
        )

        if "cgpa" not in score_type:
            continue

        try:
            values.append(
                float(
                    score.get(
                        "value",
                        ""
                    )
                )
            )
        except (
            TypeError,
            ValueError,
        ):
            continue

    return max(values) if values else None


def extract_best_percentage(
    education: list[dict[str, Any]],
) -> float | None:

    values: list[float] = []

    for item in education:

        score = item.get(
            "score",
            {}
        )

        if not isinstance(score, dict):
            continue

        score_type = normalize_key(
            score.get("type", "")
        )

        if "percentage" not in score_type:
            continue

        try:
            values.append(
                float(
                    score.get(
                        "value",
                        ""
                    )
                )
            )
        except (
            TypeError,
            ValueError,
        ):
            continue

    return max(values) if values else None


# =============================================================================
# FILE LOCATION HELPERS
# =============================================================================

def find_markdown(
    app_no: str,
) -> Path | None:

    direct = MD_FOLDER / f"{app_no}.md"

    if direct.exists():
        return direct

    if not MD_FOLDER.exists():
        return None

    for path in MD_FOLDER.glob("APP*.md"):

        if path.stem.upper() == app_no.upper():
            return path

    return None


def find_pdf(
    app_no: str,
) -> Path | None:

    direct = PDF_FOLDER / f"{app_no}.pdf"

    if direct.exists():
        return direct

    if not PDF_FOLDER.exists():
        return None

    for path in PDF_FOLDER.glob("*.pdf"):

        if path.stem.upper() == app_no.upper():
            return path

    return None


# =============================================================================
# MARKDOWN CHUNKING
# =============================================================================

def chunk_markdown(
    text: str,
    source_file: str,
) -> list[dict[str, Any]]:

    chunks: list[dict[str, Any]] = []

    current_section = "General"
    buffer: list[str] = []

    def flush():

        nonlocal buffer

        if not buffer:
            return

        content = "\n".join(
            buffer
        ).strip()

        if len(content) >= 40:

            # Add overlapping tail text to preserve context between chunks.
            if len(content) <= MD_CHUNK_LENGTH:

                chunks.append({
                    "text": content,
                    "section": current_section,
                    "source_type": "markdown",
                    "source_file": source_file,
                })

            else:

                start = 0
                while start < len(content):

                    end = start + MD_CHUNK_LENGTH

                    part = content[
                        start:end
                    ].strip()

                    if len(part) >= 40:

                        chunks.append({
                            "text": part,
                            "section": current_section,
                            "source_type": "markdown",
                            "source_file": source_file,
                        })

                    if end >= len(content):
                        break

                    start = max(
                        end - CHUNK_OVERLAP,
                        start + 1,
                    )

        buffer = []

    for raw_line in text.splitlines():

        line = raw_line.strip()

        if not line:
            continue

        if line.startswith("#"):

            flush()

            current_section = clean_text(
                line.lstrip("#")
            )

            # Keep the heading itself in the vector text so project names,
            # section names and candidate names remain searchable.
            buffer.append(current_section)

            continue

        buffer.append(line)

    flush()

    return chunks


# =============================================================================
# PDF TEXT + CHUNKING
# =============================================================================

def extract_pdf_pages(
    pdf_path: Path,
) -> list[dict[str, Any]]:

    if fitz is None:

        return []

    pages: list[dict[str, Any]] = []

    with fitz.open(
        pdf_path
    ) as pdf:

        for page_index, page in enumerate(
            pdf,
            start=1,
        ):

            text = page.get_text(
                "text"
            )

            text = clean_text(
                text
            )

            if not text:
                continue

            pages.append({
                "page": page_index,
                "text": text,
            })

    return pages


def chunk_pdf(
    pdf_path: Path,
) -> list[dict[str, Any]]:

    pages = extract_pdf_pages(
        pdf_path
    )

    chunks: list[dict[str, Any]] = []

    for page in pages:

        text = page["text"]

        start = 0

        while start < len(text):

            end = start + PDF_CHUNK_LENGTH

            part = text[
                start:end
            ].strip()

            if len(part) >= 40:

                chunks.append({
                    "text": part,
                    "section": f"PDF Page {page['page']}",
                    "source_type": "pdf",
                    "source_file": pdf_path.name,
                    "page": page["page"],
                })

            if end >= len(text):
                break

            start = max(
                end - CHUNK_OVERLAP,
                start + 1,
            )

    return chunks


# =============================================================================
# INDEX META
# =============================================================================

def file_signature(
    path: Path,
) -> dict[str, Any]:

    stat = path.stat()

    return {
        "name": path.name,
        "size": stat.st_size,
        "modified_ns": stat.st_mtime_ns,
    }


def current_source_signature() -> list[dict[str, Any]]:

    files: list[Path] = []

    if MD_FOLDER.exists():
        files.extend(
            sorted(
                MD_FOLDER.glob("APP*.md")
            )
        )

    if PDF_FOLDER.exists():
        files.extend(
            sorted(
                PDF_FOLDER.glob("APP*.pdf")
            )
        )

    return [
        file_signature(path)
        for path in files
    ]


def load_index_meta() -> dict[str, Any]:

    if not INDEX_META_PATH.exists():
        return {}

    try:

        with INDEX_META_PATH.open(
            "r",
            encoding="utf-8",
        ) as file:

            return json.load(file)

    except Exception:
        return {}


def save_index_meta(
    chunk_count: int,
) -> None:

    payload = {

        "embedding_model":
            EMBEDDING_MODEL,

        "indexed_at":
            datetime.now().isoformat(),

        "chunk_count":
            chunk_count,

        "sources":
            current_source_signature(),
    }

    with INDEX_META_PATH.open(
        "w",
        encoding="utf-8",
    ) as file:

        json.dump(
            payload,
            file,
            indent=2,
        )


def index_is_current() -> bool:

    meta = load_index_meta()

    if not meta:
        return False

    if (
        meta.get("embedding_model")
        != EMBEDDING_MODEL
    ):
        return False

    if (
        int(meta.get("chunk_count", 0))
        != collection.count()
    ):
        return False

    return (
        meta.get("sources", [])
        == current_source_signature()
    )


# =============================================================================
# BUILD VECTOR INDEX
# =============================================================================

def build_index(
    force_rebuild: bool = False,
) -> dict[str, Any]:

    if (
        not force_rebuild
        and collection.count() > 0
        and index_is_current()
    ):

        return {
            "status": "already_current",
            "chunks": collection.count(),
        }

    all_chunks: list[dict[str, Any]] = []

    # -------------------------------------------------------------------------
    # Markdown source
    # -------------------------------------------------------------------------

    if MD_FOLDER.exists():

        for md_path in sorted(
            MD_FOLDER.glob("APP*.md")
        ):

            try:

                text = md_path.read_text(
                    encoding="utf-8",
                    errors="replace",
                )

                app_no = md_path.stem.upper()

                for chunk in chunk_markdown(
                    text,
                    md_path.name,
                ):

                    chunk[
                        "application_no"
                    ] = app_no

                    all_chunks.append(
                        chunk
                    )

            except Exception:
                continue

    # -------------------------------------------------------------------------
    # Original PDF source
    # -------------------------------------------------------------------------

    if PDF_FOLDER.exists() and fitz is not None:

        for pdf_path in sorted(
            PDF_FOLDER.glob("APP*.pdf")
        ):

            try:

                app_no = pdf_path.stem.upper()

                for chunk in chunk_pdf(
                    pdf_path
                ):

                    chunk[
                        "application_no"
                    ] = app_no

                    all_chunks.append(
                        chunk
                    )

            except Exception:
                continue

    if not all_chunks:

        raise FileNotFoundError(
            "No Markdown/PDF content was found to build the retrieval index."
        )

    documents = [
        chunk["text"]
        for chunk in all_chunks
    ]

    ids = [
        (
            f"{chunk['application_no']}_"
            f"{chunk['source_type']}_"
            f"{index:05d}"
        )
        for index, chunk in enumerate(
            all_chunks
        )
    ]

    metadata = [

        {
            "application_no":
                chunk["application_no"],

            "source_type":
                chunk["source_type"],

            "source_file":
                chunk["source_file"],

            "section":
                chunk.get(
                    "section",
                    "General",
                ),

            "page":
                int(
                    chunk.get(
                        "page",
                        0,
                    )
                ),
        }

        for chunk in all_chunks
    ]

    embeddings = embedding_model.encode(
        documents,
        batch_size=64,
        show_progress_bar=True,
        normalize_embeddings=True,
    )

    # Replace old index.
    if collection.count():

        existing = collection.get()

        if existing.get("ids"):

            collection.delete(
                ids=existing["ids"]
            )

    collection.upsert(

        ids=ids,

        documents=documents,

        embeddings=embeddings.tolist(),

        metadatas=metadata,
    )

    save_index_meta(
        len(documents)
    )

    return {
        "status": "rebuilt",
        "chunks": len(documents),
        "markdown_chunks": sum(
            1
            for x in all_chunks
            if x["source_type"] == "markdown"
        ),
        "pdf_chunks": sum(
            1
            for x in all_chunks
            if x["source_type"] == "pdf"
        ),
    }


# =============================================================================
# QUERY UNDERSTANDING
# =============================================================================

def extract_json_object(
    text: str,
) -> dict[str, Any]:

    cleaned = clean_text(text)

    cleaned = re.sub(
        r"```json",
        "",
        cleaned,
        flags=re.I,
    )

    cleaned = cleaned.replace(
        "```",
        "",
    ).strip()

    match = re.search(
        r"\{.*\}",
        cleaned,
        re.DOTALL,
    )

    if not match:
        raise ValueError(
            "Query LLM did not return a JSON object."
        )

    return json.loads(
        match.group(0)
    )


def parse_query(
    question: str,
) -> dict[str, Any]:

    prompt = f"""
You are the HireNex recruiter-query understanding module.

Convert the recruiter's natural-language request into ONLY valid JSON.

Use exactly this structure:

{{
  "query_type": "structured" | "semantic" | "mixed",
  "min_experience_years": null,
  "min_experience_inclusive": true,
  "max_experience_years": null,
  "max_experience_inclusive": true,
  "skills": [],
  "degree": "",
  "college": "",
  "job_role": "",
  "certifications": [],
  "min_cgpa": null,
  "min_cgpa_inclusive": true,
  "min_percentage": null,
  "min_percentage_inclusive": true,
  "semantic_query": ""
}}

Rules:

1. "at least 2 years" => min_experience_years = 2.
2. "more than 2 years" => min_experience_years = 2.
3. "under 3 years" or "less than 3 years" => max_experience_years = 3.
4. "Python and SQL" => skills ["Python", "SQL"].
5. "CGPA above 8" => min_cgpa = 8 and min_cgpa_inclusive = false.
6. "CGPA at least 8" => min_cgpa = 8 and min_cgpa_inclusive = true.
7. "B.Tech candidates" => degree = "B.Tech".
8. "from VIT-AP" => college = "VIT-AP".
9. "for a data scientist role" => job_role = "data scientist".
10. Conceptual requests such as "strong machine learning project experience"
   should use query_type = "semantic" and put the conceptual phrase in semantic_query.
11. Requests mixing exact filters with semantic meaning should use "mixed".
12. "more than 2 years" means min_experience_years = 2 and min_experience_inclusive = false.
13. "less than 3 years" means max_experience_years = 3 and max_experience_inclusive = false.
14. "at least 2 years" means min_experience_years = 2 and min_experience_inclusive = true.
15. Do not invent values.
16. Unknown scalar fields are "" or null.
17. Unknown list fields are [].
18. semantic_query should contain only the semantic portion of the request.
19. Return JSON only.

Recruiter request:
{question}
"""

    response = client.chat.completions.create(

        model=QUERY_MODEL,

        messages=[
            {
                "role": "system",
                "content": prompt,
            }
        ],

        temperature=0,
    )

    content = (
        response
        .choices[0]
        .message
        .content
    )

    return extract_json_object(
        content
    )


# =============================================================================
# STRUCTURED MATCHING
# =============================================================================

def string_contains(
    value: str,
    query_value: str,
) -> bool:

    return (
        normalize_key(query_value)
        in normalize_key(value)
    )


def candidate_has_skill(
    candidate: dict[str, Any],
    required_skills: list[str],
) -> bool:

    candidate_skills = [
        normalize_key(x)
        for x in ensure_list(
            candidate.get("skills", [])
        )
    ]

    for required in required_skills:

        required_norm = normalize_key(
            required
        )

        if not any(
            (
                skill == required_norm
                or required_norm in skill
            )
            for skill in candidate_skills
        ):

            return False

    return True


def candidate_has_degree(
    candidate: dict[str, Any],
    requested_degree: str,
) -> bool:

    if not requested_degree:
        return True

    requested = normalize_key(
        requested_degree
    )

    for education in candidate.get(
        "education",
        [],
    ):

        degree = normalize_key(
            education.get(
                "degree",
                ""
            )
        )

        field = normalize_key(
            education.get(
                "field",
                ""
            )
        )

        if (
            requested in degree
            or requested in field
        ):
            return True

    return False


def candidate_has_college(
    candidate: dict[str, Any],
    requested_college: str,
) -> bool:

    if not requested_college:
        return True

    requested = normalize_key(
        requested_college
    )

    for education in candidate.get(
        "education",
        [],
    ):

        institute = normalize_key(
            education.get(
                "institute",
                ""
            )
        )

        if requested in institute:
            return True

    return False


def candidate_has_role(
    candidate: dict[str, Any],
    requested_role: str,
) -> bool:

    if not requested_role:
        return True

    requested = normalize_key(
        requested_role
    )

    for experience in candidate.get(
        "experience",
        [],
    ):

        role = normalize_key(
            experience.get(
                "role",
                ""
            )
        )

        if requested in role:
            return True

    return False


def candidate_has_certification(
    candidate: dict[str, Any],
    requested: list[str],
) -> bool:

    if not requested:
        return True

    candidate_certs = []

    for cert in candidate.get(
        "certifications",
        [],
    ):

        if isinstance(cert, dict):

            candidate_certs.append(
                normalize_key(
                    cert.get(
                        "name",
                        ""
                    )
                )
            )

    for required in requested:

        required_norm = normalize_key(
            required
        )

        if not any(
            required_norm in cert
            for cert in candidate_certs
        ):
            return False

    return True


def apply_structured_filters(
    candidates: list[dict[str, Any]],
    filters: dict[str, Any],
) -> list[dict[str, Any]]:

    results = []

    min_exp = filters.get(
        "min_experience_years"
    )
    min_exp_inclusive = filters.get(
        "min_experience_inclusive",
        True,
    )

    max_exp = filters.get(
        "max_experience_years"
    )
    max_exp_inclusive = filters.get(
        "max_experience_inclusive",
        True,
    )

    min_cgpa = filters.get(
        "min_cgpa"
    )

    min_percentage = filters.get(
        "min_percentage"
    )
    min_cgpa_inclusive = filters.get(
        "min_cgpa_inclusive",
        True,
    )
    min_percentage_inclusive = filters.get(
        "min_percentage_inclusive",
        True,
    )

    required_skills = [
        x
        for x in ensure_list(
            filters.get(
                "skills",
                []
            )
        )
        if clean_text(x)
    ]

    degree = clean_text(
        filters.get(
            "degree",
            ""
        )
    )

    college = clean_text(
        filters.get(
            "college",
            ""
        )
    )

    role = clean_text(
        filters.get(
            "job_role",
            ""
        )
    )

    certifications = [
        x
        for x in ensure_list(
            filters.get(
                "certifications",
                []
            )
        )
        if clean_text(x)
    ]

    for candidate in candidates:

        experience = float(
            candidate.get(
                "total_experience_years",
                0,
            )
            or 0
        )

        if min_exp is not None:

            try:

                threshold = float(min_exp)
                if min_exp_inclusive:
                    if experience < threshold:
                        continue
                else:
                    if experience <= threshold:
                        continue

            except (
                TypeError,
                ValueError,
            ):
                pass

        if max_exp is not None:

            try:

                threshold = float(max_exp)
                if max_exp_inclusive:
                    if experience > threshold:
                        continue
                else:
                    if experience >= threshold:
                        continue

            except (
                TypeError,
                ValueError,
            ):
                pass

        if required_skills:

            if not candidate_has_skill(
                candidate,
                required_skills,
            ):
                continue

        if degree:

            if not candidate_has_degree(
                candidate,
                degree,
            ):
                continue

        if college:

            if not candidate_has_college(
                candidate,
                college,
            ):
                continue

        if role:

            if not candidate_has_role(
                candidate,
                role,
            ):
                continue

        if certifications:

            if not candidate_has_certification(
                candidate,
                certifications,
            ):
                continue

        if min_cgpa is not None:

            cgpa = candidate.get(
                "max_cgpa"
            )

            try:

                if cgpa is None:
                    continue
                value = float(cgpa)
                threshold = float(min_cgpa)
                if min_cgpa_inclusive:
                    if value < threshold:
                        continue
                else:
                    if value <= threshold:
                        continue

            except (
                TypeError,
                ValueError,
            ):
                continue

        if min_percentage is not None:

            percentage = candidate.get(
                "max_percentage"
            )

            try:

                if percentage is None:
                    continue
                value = float(percentage)
                threshold = float(min_percentage)
                if min_percentage_inclusive:
                    if value < threshold:
                        continue
                else:
                    if value <= threshold:
                        continue

            except (
                TypeError,
                ValueError,
            ):
                continue

        results.append(
            candidate
        )

    return results


# =============================================================================
# SEMANTIC RETRIEVAL
# =============================================================================

def semantic_retrieve(
    question: str,
    n_results: int = DEFAULT_SEMANTIC_RESULTS,
    allowed_application_nos: set[str] | None = None,
) -> list[dict[str, Any]]:

    if collection.count() == 0:

        build_index()

    query_embedding = embedding_model.encode(
        [question],
        normalize_embeddings=True,
    )[0].tolist()

    n_results = max(
        1,
        min(
            n_results,
            collection.count(),
        ),
    )

    result = collection.query(

        query_embeddings=[
            query_embedding
        ],

        n_results=n_results,

        include=[
            "documents",
            "metadatas",
            "distances",
        ],
    )

    documents = (
        result.get("documents", [[]])[0]
    )

    metadatas = (
        result.get("metadatas", [[]])[0]
    )

    distances = (
        result.get("distances", [[]])[0]
    )

    retrieved: list[dict[str, Any]] = []

    for document, metadata, distance in zip(
        documents,
        metadatas,
        distances,
    ):

        app_no = clean_text(
            metadata.get(
                "application_no",
                ""
            )
        ).upper()

        if (
            allowed_application_nos is not None
            and app_no not in allowed_application_nos
        ):
            continue

        similarity = round(
            max(
                0.0,
                1.0 - float(distance),
            ),
            4,
        )

        retrieved.append({

            "application_no":
                app_no,

            "source_type":
                metadata.get(
                    "source_type",
                    ""
                ),

            "source_file":
                metadata.get(
                    "source_file",
                    ""
                ),

            "section":
                metadata.get(
                    "section",
                    ""
                ),

            "page":
                metadata.get(
                    "page",
                    0
                ),

            "similarity":
                similarity,

            "text":
                document,
        })

    return retrieved


def rank_candidates_by_semantics(
    retrieved_chunks: list[dict[str, Any]],
    candidates: list[dict[str, Any]],
) -> list[dict[str, Any]]:

    candidate_map = {
        candidate["application_no"]:
            candidate
        for candidate in candidates
    }

    score_map: dict[str, float] = {}
    source_types: dict[str, set[str]] = {}

    for chunk in retrieved_chunks:

        app_no = chunk[
            "application_no"
        ]

        if not app_no:
            continue

        score = float(
            chunk.get(
                "similarity",
                0,
            )
            or 0
        )

        score_map[app_no] = max(
            score_map.get(
                app_no,
                0,
            ),
            score,
        )

        source_types.setdefault(
            app_no,
            set(),
        ).add(
            chunk.get(
                "source_type",
                ""
            )
        )

    ranked = []

    for app_no, score in sorted(
        score_map.items(),
        key=lambda x: x[1],
        reverse=True,
    ):

        candidate = candidate_map.get(
            app_no
        )

        if candidate is None:
            continue

        output = candidate.copy()

        # Small bonus when both MD and original PDF independently
        # support the semantic retrieval.
        bonus = (
            0.03
            if len(
                source_types.get(
                    app_no,
                    set(),
                )
            ) >= 2
            else 0.0
        )

        output[
            "similarity_score"
        ] = round(
            min(
                score + bonus,
                1.0,
            ),
            4,
        )

        ranked.append(
            output
        )

    return ranked


# =============================================================================
# PDF EVIDENCE
# =============================================================================

def relevant_pdf_evidence(
    app_no: str,
    question: str,
    max_chars: int = 1000,
) -> str:

    pdf_path = find_pdf(
        app_no
    )

    if pdf_path is None:
        return ""

    if fitz is None:
        return ""

    pages = extract_pdf_pages(
        pdf_path
    )

    question_terms = [
        term.lower()
        for term in re.findall(
            r"[A-Za-z][A-Za-z0-9+#.-]{2,}",
            question,
        )
    ]

    for page in pages:

        text = page["text"]

        lower = text.lower()

        positions = [
            lower.find(term)
            for term in question_terms
            if lower.find(term) >= 0
        ]

        if positions:

            center = min(
                positions
            )

            start = max(
                center - max_chars // 3,
                0,
            )

            snippet = text[
                start:
                start + max_chars
            ]

            return (
                f"[PDF Page {page['page']}]\n"
                f"{snippet}"
            )

    if pages:

        return (
            f"[PDF Page {pages[0]['page']}]\n"
            f"{pages[0]['text'][:max_chars]}"
        )

    return ""


# =============================================================================
# MARKDOWN EVIDENCE
# =============================================================================

def relevant_markdown_evidence(
    app_no: str,
    question: str,
    max_chars: int = 1000,
) -> str:

    md_path = find_markdown(
        app_no
    )

    if md_path is None:
        return ""

    text = md_path.read_text(
        encoding="utf-8",
        errors="replace",
    )

    lower = text.lower()

    question_terms = [
        term.lower()
        for term in re.findall(
            r"[A-Za-z][A-Za-z0-9+#.-]{2,}",
            question,
        )
    ]

    positions = [
        lower.find(term)
        for term in question_terms
        if lower.find(term) >= 0
    ]

    if positions:

        center = min(
            positions
        )

        start = max(
            center - max_chars // 3,
            0,
        )

        return text[
            start:
            start + max_chars
        ]

    return text[
        :max_chars
    ]


# =============================================================================
# ANSWER GENERATION
# =============================================================================

def build_grounded_context(
    question: str,
    filters: dict[str, Any],
    candidates: list[dict[str, Any]],
    retrieved_chunks: list[dict[str, Any]],
) -> str:

    parts: list[str] = []

    parts.append(
        "STRUCTURED FILTERS:\n"
        + json.dumps(
            filters,
            indent=2,
            ensure_ascii=False,
        )
    )

    parts.append(
        "\nMATCHING CANDIDATES:\n"
        + json.dumps(
            {
                "total_matches": len(candidates),
                "sample_matches": [
                    {
                        "application_no": x.get("application_no"),
                        "name": x.get("name"),
                        "experience_years": x.get("total_experience_years"),
                        "skills": x.get("skills", []),
                        "cgpa": x.get("max_cgpa"),
                        "percentage": x.get("max_percentage"),
                    }
                    for x in candidates[:30]
                ],
            },
            indent=2,
            ensure_ascii=False,
        )
    )

    if retrieved_chunks:

        parts.append(
            "\nRETRIEVED DOCUMENT EVIDENCE:\n"
        )

        for index, chunk in enumerate(
            retrieved_chunks[
                :DEFAULT_CONTEXT_RESULTS
            ],
            start=1,
        ):

            source = (
                f"{chunk.get('source_type', '')}"
                f"::{chunk.get('source_file', '')}"
            )

            if chunk.get("page"):
                source += (
                    f"::page {chunk['page']}"
                )

            parts.append(
                f"[Evidence {index} | {source}]\n"
                f"{chunk.get('text', '')}"
            )

    # Add original PDF evidence for a few candidates.
    for candidate in candidates[:5]:

        app_no = candidate.get(
            "application_no",
            "",
        )

        pdf_evidence = (
            relevant_pdf_evidence(
                app_no,
                question,
                800,
            )
        )

        if pdf_evidence:

            parts.append(
                "\nORIGINAL PDF EVIDENCE "
                f"FOR {app_no}:\n"
                f"{pdf_evidence}"
            )

    return "\n\n".join(
        parts
    )


def ask_rag(
    question: str,
    verbose: bool = False,
) -> dict[str, Any]:

    candidates = merge_candidates(
        *load_source_jsons()
    )

    if not candidates:

        return {
            "answer":
                "No parsed candidate data is available. "
                "Run NLP and/or LLM processing first.",
            "filters": {},
            "candidates": [],
            "retrieved_chunks": [],
            "excel_path": None,
        }

    filters = parse_query(
        question
    )

    query_type = filters.get(
        "query_type",
        "structured",
    )

    retrieved_chunks: list[
        dict[str, Any]
    ] = []

    filtered_candidates = candidates

    # -------------------------------------------------------------------------
    # Structured path
    # -------------------------------------------------------------------------

    if query_type in {
        "structured",
        "mixed",
    }:

        filtered_candidates = (
            apply_structured_filters(
                candidates,
                filters,
            )
        )

    # -------------------------------------------------------------------------
    # Semantic path
    # -------------------------------------------------------------------------

    if query_type in {
        "semantic",
        "mixed",
    }:

        semantic_query = clean_text(
            filters.get(
                "semantic_query",
                ""
            )
        ) or question

        allowed = None

        if query_type == "mixed":

            allowed = {
                x["application_no"]
                for x in filtered_candidates
            }

        retrieved_chunks = semantic_retrieve(
            semantic_query,
            DEFAULT_SEMANTIC_RESULTS,
            allowed,
        )

        if query_type == "semantic":

            filtered_candidates = (
                rank_candidates_by_semantics(
                    retrieved_chunks,
                    candidates,
                )
            )

        else:

            ranked = (
                rank_candidates_by_semantics(
                    retrieved_chunks,
                    filtered_candidates,
                )
            )

            # Keep deterministic filter results even if
            # the semantic query is not strong enough to rank all.
            ranked_map = {
                x["application_no"]: x
                for x in ranked
            }

            filtered_candidates = [
                ranked_map.get(
                    x["application_no"],
                    x,
                )
                for x in filtered_candidates
            ]

            filtered_candidates.sort(
                key=lambda x:
                    float(
                        x.get(
                            "similarity_score",
                            0,
                        )
                        or 0
                    ),
                reverse=True,
            )

    # -------------------------------------------------------------------------
    # Grounded answer
    # -------------------------------------------------------------------------

    context = build_grounded_context(
        question,
        filters,
        filtered_candidates,
        retrieved_chunks,
    )

    system_prompt = """
You are HireNex, an AI recruitment retrieval assistant.

Answer the recruiter's question using ONLY the supplied structured
candidate data and retrieved document evidence.

Rules:
1. Never invent a candidate, skill, experience, company, date, or score.
2. For exact filters, trust the deterministic matching results.
3. Do not override Python filtering decisions.
4. When no candidates match, clearly say that no matching candidates were found.
5. Mention application numbers when referring to candidates.
6. Keep the response concise and useful to a recruiter.
7. If evidence is insufficient, say that the available data is insufficient.
"""

    user_prompt = (
        f"QUESTION:\n{question}\n\n"
        f"CONTEXT:\n{context}"
    )

    response = client.chat.completions.create(

        model=ANSWER_MODEL,

        messages=[
            {
                "role": "system",
                "content": system_prompt,
            },
            {
                "role": "user",
                "content": user_prompt,
            },
        ],

        temperature=0.1,
    )

    answer = (
        response
        .choices[0]
        .message
        .content
        .strip()
    )

    excel_path = None

    if filtered_candidates:

        excel_path = create_results_excel(
            question,
            filters,
            filtered_candidates,
        )

    if verbose:

        print("\n" + "=" * 80)
        print(f"Question: {question}")
        print("-" * 80)
        print(
            json.dumps(
                filters,
                indent=2,
                ensure_ascii=False,
            )
        )
        print("-" * 80)

        for candidate in filtered_candidates:

            score = candidate.get(
                "similarity_score"
            )

            score_text = (
                f" | similarity={score}"
                if score is not None
                else ""
            )

            print(
                f"{candidate.get('application_no')} | "
                f"{candidate.get('name')} | "
                f"{candidate.get('total_experience_years')} years"
                f"{score_text}"
            )

        print("-" * 80)
        print(answer)
        print("=" * 80)

    return {
        "answer": answer,
        "filters": filters,
        "candidates": filtered_candidates,
        "retrieved_chunks": retrieved_chunks,
        "excel_path": excel_path,
    }


# =============================================================================
# EXCEL GENERATION
# =============================================================================

def create_results_excel(
    question: str,
    filters: dict[str, Any],
    candidates: list[dict[str, Any]],
) -> str:

    if Workbook is None:
        raise ImportError(
            "openpyxl is required for Excel generation. "
            "Run: python -m pip install openpyxl"
        )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    output_path = (
        RAG_OUTPUT_FOLDER
        / f"HireNex_Results_{timestamp}.xlsx"
    )

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Candidates"

    headers = [
        "Application_No",
        "Name",
        "Experience_Years",
        "Skills",
        "Email",
        "Phone",
        "LinkedIn",
        "GitHub",
        "Portfolio",
        "CGPA",
        "Percentage",
        "Education",
        "Experience",
        "Projects",
        "Certifications",
    ]

    sheet.append(headers)

    for candidate in candidates:
        sheet.append([
            candidate.get("application_no", ""),
            candidate.get("name", ""),
            float(candidate.get("total_experience_years", 0) or 0),
            ", ".join(candidate.get("skills", [])),
            candidate.get("email", ""),
            candidate.get("phone", ""),
            candidate.get("linkedin", ""),
            candidate.get("github", ""),
            candidate.get("portfolio", ""),
            candidate.get("max_cgpa", ""),
            candidate.get("max_percentage", ""),
            json.dumps(candidate.get("education", []), ensure_ascii=False),
            json.dumps(candidate.get("experience", []), ensure_ascii=False),
            json.dumps(candidate.get("projects", []), ensure_ascii=False),
            json.dumps(candidate.get("certifications", []), ensure_ascii=False),
        ])

    # Professional, readable formatting.
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="111827",
    )
    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row in sheet.iter_rows(
        min_row=2,
        max_row=sheet.max_row,
        min_col=1,
        max_col=sheet.max_column,
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    # Number formats.
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=3).number_format = "0.0"
        sheet.cell(row=row, column=10).number_format = "0.00"
        sheet.cell(row=row, column=11).number_format = "0.00"

    widths = {
        "A": 18,
        "B": 28,
        "C": 18,
        "D": 42,
        "E": 32,
        "F": 18,
        "G": 34,
        "H": 34,
        "I": 34,
        "J": 12,
        "K": 14,
        "L": 45,
        "M": 55,
        "N": 55,
        "O": 40,
    }

    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.freeze_panes = "A2"

    if sheet.max_row > 1:
        table_ref = f"A1:O{sheet.max_row}"
        table = Table(
            displayName="HireNexCandidates",
            ref=table_ref,
        )
        table_style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        table.tableStyleInfo = table_style
        sheet.add_table(table)

    # Query details sheet.
    query_sheet = workbook.create_sheet("Query")
    query_sheet.append(["Field", "Value"])
    query_sheet.append(["Recruiter_Query", question])
    query_sheet.append([
        "Filters",
        json.dumps(filters, ensure_ascii=False),
    ])
    query_sheet.append([
        "Generated_At",
        datetime.now().isoformat(),
    ])

    for cell in query_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in query_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    query_sheet.column_dimensions["A"].width = 24
    query_sheet.column_dimensions["B"].width = 90
    query_sheet.freeze_panes = "A2"

    workbook.save(
        output_path
    )

    return str(output_path)


def clear_retrieval_index() -> int:
    """Remove all indexed retrieval chunks while keeping the Chroma database reusable."""
    global collection
    removed = 0

    try:
        existing = collection.get()
        ids = existing.get("ids") or []
        removed = len(ids)
        if ids:
            collection.delete(ids=ids)
    except Exception:
        # Try recreating the collection if the existing handle is stale.
        try:
            chroma_client.delete_collection("hinenex_resume_knowledge")
        except Exception:
            pass
        collection = chroma_client.get_or_create_collection(
            name="hinenex_resume_knowledge",
            metadata={"hnsw:space": "cosine"},
        )

    if INDEX_META_PATH.exists():
        try:
            INDEX_META_PATH.unlink()
        except FileNotFoundError:
            pass

    return removed


# =============================================================================
# APP-FRIENDLY SEARCH FUNCTION
# =============================================================================

def search_candidates(
    question: str,
) -> dict[str, Any]:

    """
    Main function intended for app.py.

    Returns:
        {
            "answer": str,
            "filters": dict,
            "candidates": list[dict],
            "retrieved_chunks": list[dict],
            "excel_path": str | None,
        }
    """

    return ask_rag(
        question,
        verbose=False,
    )


# =============================================================================
# STATUS
# =============================================================================

def get_status() -> dict[str, Any]:

    nlp_count = len(
        load_json_file(
            NLP_JSON_PATH
        )
    )

    llm_count = len(
        load_json_file(
            LLM_JSON_PATH
        )
    )

    md_count = (
        len(
            list(
                MD_FOLDER.glob("APP*.md")
            )
        )
        if MD_FOLDER.exists()
        else 0
    )

    pdf_count = (
        len(
            list(
                PDF_FOLDER.glob("APP*.pdf")
            )
        )
        if PDF_FOLDER.exists()
        else 0
    )

    return {

        "nlp_records":
            nlp_count,

        "llm_records":
            llm_count,

        "markdown_files":
            md_count,

        "pdf_files":
            pdf_count,

        "vector_chunks":
            collection.count(),

        "index_current":
            index_is_current(),

        "nlp_json":
            str(NLP_JSON_PATH),

        "llm_json":
            str(LLM_JSON_PATH),

        "markdown_folder":
            str(MD_FOLDER),

        "pdf_folder":
            str(PDF_FOLDER),

        "rag_output":
            str(RAG_OUTPUT_FOLDER),
    }


# =============================================================================
# CLI TEST MODE
# =============================================================================

if __name__ == "__main__":

    print("HireNex Retrieval System")
    print("=" * 60)

    print(
        json.dumps(
            get_status(),
            indent=2,
            ensure_ascii=False,
        )
    )

    if collection.count() == 0:

        print(
            "\nBuilding initial vector index..."
        )

        print(
            json.dumps(
                build_index(),
                indent=2,
            )
        )

    question = input(
        "\nAsk HireNex: "
    ).strip()

    if question:

        result = ask_rag(
            question,
            verbose=True,
        )

        print(
            "\nExcel:"
            f" {result.get('excel_path')}"
        )
